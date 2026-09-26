"""P2.2.0: the reference script reproduces expected/ exactly; week 3 has drifted columns."""
import filecmp
import shutil
import subprocess
import sys
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[4]
UC1 = ROOT / "data" / "artifacts" / "uc1"


def test_reference_script_reproduces_expected_files_exactly(tmp_path):
    copy = tmp_path / "uc1"
    shutil.copytree(UC1, copy)
    before = tmp_path / "expected_before"
    shutil.copytree(copy / "expected", before)
    proc = subprocess.run([sys.executable, "reference.py"], cwd=copy, capture_output=True, text=True, check=False)
    assert proc.returncode == 0, proc.stderr
    for path in sorted(before.iterdir()):
        assert filecmp.cmp(path, copy / "expected" / path.name, shallow=False), path.name
    assert len(list(before.iterdir())) == 8


def test_columns_and_mess():
    assert list(pd.read_excel(UC1 / "week1.xlsx").columns) == ["date", "reg", "product", "amt"]
    w3 = pd.read_excel(UC1 / "week3.xlsx")
    assert list(w3.columns) == ["date", "region_name", "product", "amount"]
    assert w3["amount"].isna().any() and w3["region_name"].isna().any()
    from openpyxl import load_workbook

    cells = [r[3] for r in load_workbook(UC1 / "week3.xlsx").active.iter_rows(min_row=2, values_only=True)]
    assert any(isinstance(v, str) for v in cells)  # numbers stored as text in the xlsx
