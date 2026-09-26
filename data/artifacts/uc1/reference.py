from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parent
EXPECTED = ROOT / "expected"

WEEKS = {
    1: [
        ("North", "Hardware", 1240.50),
        ("South", "Hardware", 980.00),
        ("North", "Software", 2225.25),
        ("West", "Services", 870.00),
        ("South", "Software", 1410.75),
    ],
    2: [
        ("North", "Hardware", 1325.00),
        ("South", "Services", 1112.40),
        ("East", "Software", 1750.00),
        ("West", "Hardware", 920.35),
        ("East", "Services", 615.20),
    ],
    3: [
        ("North", "Hardware", 1420.00),
        ("South", "Hardware", 1015.75),
        ("East", "Software", 1999.99),
        ("West", "Services", 1201.10),
        ("North", "Services", 740.00),
    ],
    4: [
        ("North", "Hardware", 1560.10),
        ("South", "Software", 1675.00),
        ("East", "Services", 830.80),
        ("West", "Hardware", 1105.25),
        ("West", "Software", 1290.45),
    ],
}

HEADERS = {
    1: ("Region", "Category", "Revenue"),
    2: ("Region", "Category", "Revenue"),
    3: ("Sales Region", "Business Line", "Net Revenue"),
    4: ("Region", "Category", "Revenue"),
}

COLUMN_ALIASES = {
    "Sales Region": "Region",
    "Business Line": "Category",
    "Net Revenue": "Revenue",
}


def write_workbooks() -> None:
    ROOT.mkdir(parents=True, exist_ok=True)
    for week, rows in WEEKS.items():
        wb = Workbook()
        ws = wb.active
        ws.title = "Sales"
        ws.append(HEADERS[week])
        for row in rows:
            ws.append(row)
        wb.save(ROOT / f"week{week}.xlsx")


def load_sales(path: Path) -> list[dict[str, str]]:
    wb = load_workbook(path, data_only=True)
    ws = wb["Sales"]
    headers = [COLUMN_ALIASES.get(cell.value, cell.value) for cell in ws[1]]
    rows: list[dict[str, str]] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values, strict=True))
        rows.append(
            {
                "Region": str(row["Region"]),
                "Category": str(row["Category"]),
                "Revenue": f"{float(row['Revenue']):.2f}",
            }
        )
    return rows


def pivot(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    totals: dict[tuple[str, str], float] = {}
    for row in rows:
        key = (row["Region"], row["Category"])
        totals[key] = totals.get(key, 0.0) + float(row["Revenue"])
    return [
        {"Region": region, "Category": category, "Revenue": f"{revenue:.2f}"}
        for (region, category), revenue in sorted(totals.items())
    ]


def chart(pivot_rows: list[dict[str, str]]) -> dict:
    by_region: dict[str, float] = {}
    for row in pivot_rows:
        by_region[row["Region"]] = by_region.get(row["Region"], 0.0) + float(row["Revenue"])
    return {
        "kind": "bar",
        "title": "Revenue by Region",
        "x": "Region",
        "y": "Revenue",
        "series": [
            {"region": region, "revenue": round(revenue, 2)}
            for region, revenue in sorted(by_region.items())
        ],
    }


def write_expected() -> None:
    EXPECTED.mkdir(parents=True, exist_ok=True)
    for week in sorted(WEEKS):
        rows = pivot(load_sales(ROOT / f"week{week}.xlsx"))
        with (EXPECTED / f"week{week}_pivot.csv").open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=["Region", "Category", "Revenue"])
            writer.writeheader()
            writer.writerows(rows)
        (EXPECTED / f"week{week}_chart.json").write_text(
            json.dumps(chart(rows), indent=2, sort_keys=True) + "\n",
            encoding="utf-8",
        )


def main() -> None:
    write_workbooks()
    write_expected()


if __name__ == "__main__":
    main()
